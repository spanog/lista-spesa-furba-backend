"""Build the versioned municipality seed migration from official ISTAT files.

Usage:
  python -m scripts.build_municipalities_migration \
    --codes /path/Elenco-comuni-italiani.xlsx \
    --boundaries /path/Limiti01012026_g.zip \
    --output supabase/migrations/<timestamp>_municipalities.sql
"""
from __future__ import annotations

import argparse
import math
import struct
import unicodedata
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook


@dataclass(frozen=True)
class Municipality:
    code: str
    name: str
    province_name: str
    province_code: str
    lat: float
    lng: float


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--codes", type=Path, required=True)
    parser.add_argument("--boundaries", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    return parser.parse_args()


def normalize(value: str) -> str:
    decomposed = unicodedata.normalize("NFD", value)
    plain = "".join(char for char in decomposed if not unicodedata.combining(char))
    return " ".join(plain.casefold().split())


def sql_literal(value: str) -> str:
    return "'" + value.replace("'", "''") + "'"


def read_codes(path: Path) -> dict[str, tuple[str, str, str]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    next(rows)
    return {
        str(row[4]).zfill(6): (str(row[6]), str(row[11]), str(row[14]))
        for row in rows
        if row[4] is not None
    }


def dbf_fields(handle) -> tuple[int, int, list[tuple[str, int]]]:
    header = handle.read(32)
    count, header_length, record_length = struct.unpack("<IHH", header[4:12])
    fields = []
    field = handle.read(32)
    while field[0] != 13:
        fields.append((field[:11].split(b"\0")[0].decode(), field[16]))
        field = handle.read(32)
    return count, header_length, fields


def read_dbf(path: Path) -> list[dict[str, str]]:
    with path.open("rb") as handle:
        count, header_length, fields = dbf_fields(handle)
        handle.seek(header_length)
        record_length = 1 + sum(width for _, width in fields)
        return [parse_dbf_record(handle.read(record_length), fields) for _ in range(count)]


def parse_dbf_record(record: bytes, fields: list[tuple[str, int]]) -> dict[str, str]:
    offset = 1
    row = {}
    for name, width in fields:
        row[name] = record[offset : offset + width].decode("utf-8").strip()
        offset += width
    return row


def ring_centroid(points: list[tuple[float, float]]) -> tuple[float, float, float]:
    area2 = sum(
        x1 * y2 - x2 * y1
        for (x1, y1), (x2, y2) in zip(points, points[1:] + points[:1])
    )
    if not area2:
        return 0.0, points[0][0], points[0][1]
    edges = zip(points, points[1:] + points[:1])
    cx = sum((x1 + x2) * (x1 * y2 - x2 * y1) for (x1, y1), (x2, y2) in edges)
    edges = zip(points, points[1:] + points[:1])
    cy = sum((y1 + y2) * (x1 * y2 - x2 * y1) for (x1, y1), (x2, y2) in edges)
    return area2, cx / (3 * area2), cy / (3 * area2)


def polygon_centroid(content: bytes) -> tuple[float, float]:
    parts_count, points_count = struct.unpack("<2i", content[36:44])
    parts = list(struct.unpack(f"<{parts_count}i", content[44 : 44 + parts_count * 4]))
    point_offset = 44 + parts_count * 4
    points = list(struct.iter_unpack("<2d", content[point_offset : point_offset + points_count * 16]))
    rings = [points[start : parts[index + 1] if index + 1 < parts_count else points_count] for index, start in enumerate(parts)]
    centroids = [ring_centroid(ring) for ring in rings if len(ring) > 2]
    total_area = sum(area for area, _, _ in centroids)
    if total_area:
        x = sum(area * cx for area, cx, _ in centroids) / total_area
        y = sum(area * cy for area, _, cy in centroids) / total_area
        return x, y
    return struct.unpack("<2d", content[4:20])


def footpoint_latitude(meridional_arc: float) -> float:
    eccentricity_sq = 0.0066943799901413165
    a = 6378137.0
    denominator = a * (
        1 - eccentricity_sq / 4 - 3 * eccentricity_sq**2 / 64
        - 5 * eccentricity_sq**3 / 256
    )
    mu = meridional_arc / denominator
    e1 = (1 - math.sqrt(1 - eccentricity_sq)) / (1 + math.sqrt(1 - eccentricity_sq))
    return (
        mu + (3 * e1 / 2 - 27 * e1**3 / 32) * math.sin(2 * mu)
        + (21 * e1**2 / 16 - 55 * e1**4 / 32) * math.sin(4 * mu)
        + (151 * e1**3 / 96) * math.sin(6 * mu)
        + (1097 * e1**4 / 512) * math.sin(8 * mu)
    )


def utm32_to_wgs84(easting: float, northing: float) -> tuple[float, float]:
    a, eccentricity_sq, scale = 6378137.0, 0.0066943799901413165, 0.9996
    latitude_1 = footpoint_latitude(northing / scale)
    eccentricity_prime_sq = eccentricity_sq / (1 - eccentricity_sq)
    cos_latitude = math.cos(latitude_1)
    tangent_sq = math.tan(latitude_1) ** 2
    n1 = a / math.sqrt(1 - eccentricity_sq * math.sin(latitude_1) ** 2)
    r1 = a * (1 - eccentricity_sq) / (
        1 - eccentricity_sq * math.sin(latitude_1) ** 2
    ) ** 1.5
    d = (easting - 500000.0) / (n1 * scale)
    latitude = latitude_1 - _latitude_correction(
        d, n1, r1, tangent_sq, eccentricity_prime_sq, latitude_1
    )
    longitude = math.radians(9) + _longitude_correction(
        d, tangent_sq, eccentricity_prime_sq, cos_latitude
    )
    return math.degrees(latitude), math.degrees(longitude)


def _latitude_correction(d, n1, r1, tangent_sq, eccentricity_prime_sq, latitude):
    first = d**2 / 2
    second = (5 + 3 * tangent_sq + 10 * eccentricity_prime_sq - 4 * eccentricity_prime_sq**2 - 9 * eccentricity_prime_sq) * d**4 / 24
    third = (61 + 90 * tangent_sq + 298 * eccentricity_prime_sq + 45 * tangent_sq**2 - 252 * eccentricity_prime_sq - 3 * eccentricity_prime_sq**2) * d**6 / 720
    return n1 * math.tan(latitude) * (first - second + third) / r1


def _longitude_correction(d, tangent_sq, eccentricity_prime_sq, cosine):
    first = d - (1 + 2 * tangent_sq + eccentricity_prime_sq) * d**3 / 6
    second = (5 - 2 * eccentricity_prime_sq + 28 * tangent_sq - 3 * eccentricity_prime_sq**2 + 8 * eccentricity_prime_sq + 24 * tangent_sq**2) * d**5 / 120
    return (first + second) / cosine


def read_centers(path: Path) -> dict[str, tuple[float, float]]:
    with path.open("rb") as handle:
        handle.seek(100)
        rows = []
        while header := handle.read(8):
            _, length_words = struct.unpack(">2i", header)
            content = handle.read(length_words * 2)
            if struct.unpack("<i", content[:4])[0] in {5, 15, 25}:
                rows.append(polygon_centroid(content))
    return {
        code: utm32_to_wgs84(easting, northing)
        for code, (easting, northing) in zip(
            (row["PRO_COM_T"] for row in read_dbf(path.with_suffix(".dbf"))),
            rows,
            strict=True,
        )
    }


def merged_center(centers: dict[str, tuple[float, float]]) -> tuple[float, float]:
    points = [centers[code] for code in ("024027", "024071")]
    return tuple(sum(point[index] for point in points) / len(points) for index in range(2))


def municipalities(codes: dict[str, tuple[str, str, str]], centers: dict[str, tuple[float, float]]) -> Iterable[Municipality]:
    centers["024129"] = merged_center(centers)
    for code, (name, province_name, province_code) in sorted(codes.items()):
        lat, lng = centers[code]
        yield Municipality(code, name, province_name, province_code, lat, lng)


def municipality_values(rows: Iterable[Municipality]) -> str:
    return ",\n".join(_municipality_value(row) for row in rows)


def _municipality_value(row: Municipality) -> str:
    values = (
        sql_literal(row.code),
        sql_literal(row.name),
        sql_literal(row.province_name),
        sql_literal(row.province_code),
        sql_literal(normalize(row.name)),
    )
    center = f"ST_SetSRID(ST_MakePoint({row.lng:.7f}, {row.lat:.7f}), 4326)::extensions.geography"
    return "  (" + ", ".join((*values, center)) + ")"


def migration_sql(rows: Iterable[Municipality]) -> str:
    return """-- Source: ISTAT Elenco Comuni 2026-02-21 and Limiti amministrativi 2026-01-01.
-- Generated by scripts.build_municipalities_migration; do not hand-edit seed rows.
CREATE TABLE public.municipalities (
  code TEXT PRIMARY KEY CHECK (code ~ '^[0-9]{6}$'),
  name TEXT NOT NULL,
  province_name TEXT NOT NULL,
  province_code TEXT NOT NULL,
  normalized_name TEXT NOT NULL,
  center extensions.geography(Point, 4326) NOT NULL
);

ALTER TABLE public.municipalities ENABLE ROW LEVEL SECURITY;

CREATE POLICY municipalities_service_read ON public.municipalities
  FOR SELECT TO service_role USING (true);

INSERT INTO public.municipalities (
  code, name, province_name, province_code, normalized_name, center
) VALUES
""" + municipality_values(rows) + ";\n"


def main() -> None:
    args = parse_args()
    with TemporaryDirectory() as directory:
        with zipfile.ZipFile(args.boundaries) as archive:
            archive.extractall(directory)
        shape = next(Path(directory).glob("Com*_g/*_WGS84.shp"))
        rows = municipalities(read_codes(args.codes), read_centers(shape))
        args.output.write_text(migration_sql(rows), encoding="utf-8")


if __name__ == "__main__":
    main()
